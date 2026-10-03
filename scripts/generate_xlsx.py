#!/usr/bin/env python3
"""
SCSSA Excel Report Generator
Reads the module/batch-specific CSV from the private docs repo (DOCS_REPO env var),
builds a formatted .xlsx file, and commits it back to the same docs repo.
This means the admin can always download the latest Excel directly from GitHub.
"""

import base64
import csv
import io
import json
import os
import re
import sys
import urllib.error
import urllib.request

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Error: openpyxl is not installed. Run: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

CSV_PATH = "FSSD-24-25-Projects.csv"
XLSX_PATH = "FSSD-24-25-Projects.xlsx"

# Colour palette
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
EVEN_FILL   = PatternFill("solid", fgColor="D6E4F0")
ODD_FILL    = PatternFill("solid", fgColor="EBF5FB")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT   = Font(name="Calibri", size=10)
LINK_FONT   = Font(name="Calibri", size=10, color="0563C1", underline="single")
THIN        = Side(style="thin", color="B0BEC5")
BORDER      = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Consolidated Excel columns (1 row per group, members stacked in cells)
EXCEL_HEADERS = [
    "Group No", "Repo Name", "Repo Link", "Project Title",
    "Module", "Academic Year",
    "Student Numbers", "Names", "GitHub Usernames",
]

COL_WIDTHS = {
    "Group No": 10, "Repo Name": 38, "Repo Link": 52,
    "Project Title": 22, "Module": 10, "Academic Year": 15,
    "Student Numbers": 20, "Names": 24, "GitHub Usernames": 22,
}

MEMBER_FIELDS = {
    "Student Numbers":  ["Member1 Student No", "Member2 Student No", "Member3 Student No", "Member4 Student No"],
    "Names":            ["Member1 Name",        "Member2 Name",        "Member3 Name",        "Member4 Name"],
    "GitHub Usernames": ["Member1 GitHub",       "Member2 GitHub",      "Member3 GitHub",      "Member4 GitHub"],
}


def api_request(method: str, endpoint: str, token: str, data: dict = None):
    url = f"https://api.github.com/{endpoint.lstrip('/')}"
    headers = {
        "Accept": "application/vnd.github+json",
        "User-Agent": "scssa-xlsx-generator",
        "Authorization": f"Bearer {token}",
    }
    body = json.dumps(data).encode("utf-8") if data else None
    if body:
        headers["Content-Type"] = "application/json"
    req = urllib.request.Request(url, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req) as resp:
            content = resp.read().decode("utf-8")
            return resp.status, json.loads(content) if content else {}
    except urllib.error.HTTPError as e:
        return e.code, {}
    except Exception as e:
        return 500, {"message": str(e)}


def fetch_csv(docs_repo: str, token: str):
    status, resp = api_request("GET", f"/repos/{docs_repo}/contents/{CSV_PATH}", token)
    if status == 404:
        print(f"CSV not found in '{docs_repo}'. No Excel to generate.", file=sys.stderr)
        sys.exit(0)
    if status != 200:
        print(f"Error fetching CSV ({status}): {resp}", file=sys.stderr)
        sys.exit(1)
    raw = base64.b64decode(resp["content"].replace("\n", "")).decode("utf-8")
    return list(csv.DictReader(io.StringIO(raw)))


def consolidate_row(csv_row: dict) -> dict:
    def stack(fields):
        return "\n".join(csv_row.get(f, "").strip() for f in fields if csv_row.get(f, "").strip())
    return {
        "Group No":         csv_row.get("Group No", ""),
        "Repo Name":        csv_row.get("Repo Name", ""),
        "Repo Link":        csv_row.get("Repo Link", ""),
        "Project Title":    csv_row.get("Project Title", ""),
        "Module":           csv_row.get("Module", ""),
        "Academic Year":    csv_row.get("Academic Year", ""),
        "Student Numbers":  stack(MEMBER_FIELDS["Student Numbers"]),
        "Names":            stack(MEMBER_FIELDS["Names"]),
        "GitHub Usernames": stack(MEMBER_FIELDS["GitHub Usernames"]),
    }


def build_xlsx(csv_rows: list) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Project Proposals"
    ws.freeze_panes = "A2"

    # Header row
    for col_idx, col_name in enumerate(EXCEL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.border    = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    # Data rows
    for row_idx, csv_row in enumerate(csv_rows, start=2):
        row  = consolidate_row(csv_row)
        fill = EVEN_FILL if row_idx % 2 == 0 else ODD_FILL
        lines = row["Student Numbers"].count("\n") + 1
        ws.row_dimensions[row_idx].height = max(18, lines * 18)

        for col_idx, col_name in enumerate(EXCEL_HEADERS, start=1):
            value = row.get(col_name, "")
            if col_name == "Repo Link" and value.startswith("http"):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.hyperlink = value
                cell.font = LINK_FONT
            else:
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = BODY_FONT
            cell.fill   = fill
            cell.border = BORDER
            if col_name in ("Student Numbers", "Names", "GitHub Usernames"):
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=False)

    for col_idx, col_name in enumerate(EXCEL_HEADERS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = COL_WIDTHS.get(col_name, 15)
    ws.auto_filter.ref = ws.dimensions
    return wb


def commit_xlsx(docs_repo: str, token: str, wb: openpyxl.Workbook):
    """Commits the xlsx workbook directly to the docs repo via GitHub Contents API."""
    endpoint = f"/repos/{docs_repo}/contents/{XLSX_PATH}"

    # Get existing SHA if file already exists (needed for update)
    status, resp = api_request("GET", endpoint, token)
    file_sha = resp.get("sha") if status == 200 else None

    # Serialize xlsx to bytes in memory
    buf = io.BytesIO()
    wb.save(buf)
    encoded = base64.b64encode(buf.getvalue()).decode("utf-8")

    commit_data = {
        "message": f"chore: regenerate {XLSX_PATH} [skip ci]",
        "content": encoded,
        "committer": {
            "name": "SCSSA Bot",
            "email": "bot@scssa-uok.github.io",
        },
    }
    if file_sha:
        commit_data["sha"] = file_sha

    put_status, put_resp = api_request("PUT", endpoint, token, commit_data)
    if put_status in (200, 201):
        xlsx_url = f"https://github.com/{docs_repo}/blob/main/{XLSX_PATH}"
        print(f"Excel report committed to '{docs_repo}/{XLSX_PATH}'")
        print(f"Download: {xlsx_url}")
    else:
        print(f"Warning: Failed to commit xlsx ({put_status}): {put_resp}", file=sys.stderr)


def main():
    token     = os.environ.get("APP_TOKEN") or os.environ.get("ISSUE_TOKEN") or os.environ.get("GITHUB_TOKEN", "")
    docs_repo = os.environ.get("DOCS_REPO") or os.environ.get("REPO_FULL_NAME", "")
    record_basename = os.environ.get("PROJECT_RECORD_BASENAME")
    if record_basename:
        global CSV_PATH, XLSX_PATH
        CSV_PATH = f"{record_basename}.csv"
        XLSX_PATH = f"{record_basename}.xlsx"

    if not token or not docs_repo:
        print("Error: APP_TOKEN and DOCS_REPO must be set.", file=sys.stderr)
        sys.exit(1)

    print(f"Fetching '{CSV_PATH}' from '{docs_repo}'...")
    rows = fetch_csv(docs_repo, token)
    print(f"Building Excel workbook ({len(rows)} group(s))...")
    wb = build_xlsx(rows)
    print(f"Committing '{XLSX_PATH}' to '{docs_repo}'...")
    commit_xlsx(docs_repo, token, wb)


if __name__ == "__main__":
    main()
